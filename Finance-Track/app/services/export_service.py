"""Export service for generating PDF and Excel reports."""
import io
import csv
from datetime import date


def export_csv(data, headers, filename='export.csv'):
    """Export data as CSV."""
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(headers)
    for row in data:
        writer.writerow(row)
    output.seek(0)
    return output.getvalue()


def export_excel(data, headers, filename='export.xlsx'):
    """Export data as Excel."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

        wb = Workbook()
        ws = wb.active
        ws.title = 'Report'

        header_font = Font(bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='4263EB', end_color='4263EB', fill_type='solid')

        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center')

        for row_idx, row in enumerate(data, 2):
            for col_idx, value in enumerate(row, 1):
                ws.cell(row=row_idx, column=col_idx, value=value)

        for col in ws.columns:
            max_length = 0
            for cell in col:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            ws.column_dimensions[col[0].column_letter].width = min(max_length + 2, 30)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output.getvalue()
    except ImportError:
        return None


def export_pdf(data, headers, title='Report', filename='export.pdf'):
    """Export data as PDF."""
    try:
        from reportlab.lib.pagesizes import letter, landscape
        from reportlab.lib import colors
        from reportlab.lib.units import inch
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

        output = io.BytesIO()
        doc = SimpleDocTemplate(output, pagesize=landscape(letter))
        styles = getSampleStyleSheet()
        elements = []

        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=18,
            spaceAfter=20,
        )
        elements.append(Paragraph(title, title_style))
        elements.append(Spacer(1, 12))

        table_data = [headers] + data
        table = Table(table_data)
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#4263EB')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, 0), 'DejaVuSans-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 10),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
            ('BACKGROUND', (0, 1), (-1, -1), colors.white),
            ('GRID', (0, 0), (-1, -1), 1, colors.grey),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')]),
        ]))
        elements.append(table)

        doc.build(elements)
        output.seek(0)
        return output.getvalue()
    except ImportError:
        return None


def generate_pdf_report(report_data):
    """Generate a comprehensive PDF finance report.

    Args:
        report_data: dict containing:
            - date_from: str (YYYY-MM-DD)
            - date_to: str (YYYY-MM-DD)
            - user_name: str
            - total_income: float
            - total_expenses: float
            - net_savings: float
            - savings_rate: float
            - income_count: int
            - expense_count: int
            - category_breakdown: list of dicts (category_name, total, count, percentage)
            - transactions: list of dicts (date, type, category, notes, amount)
            - payment_methods: list of dicts (method, total, count, percentage)
    """
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib import colors
        from reportlab.lib.units import mm
        from reportlab.platypus import (
            SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, PageBreak
        )
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont

        # Register DejaVu Sans font for Unicode support (₹ symbol)
        pdfmetrics.registerFont(TTFont('DejaVuSans', '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'))
        pdfmetrics.registerFont(TTFont('DejaVuSans-Bold', '/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf'))

        output = io.BytesIO()
        page_width, page_height = A4
        margin = 20 * mm

        doc = SimpleDocTemplate(
            output,
            pagesize=A4,
            leftMargin=margin,
            rightMargin=margin,
            topMargin=margin,
            bottomMargin=margin
        )

        styles = getSampleStyleSheet()
        elements = []

        # Custom styles with DejaVu Sans font
        title_style = ParagraphStyle(
            'ReportTitle',
            parent=styles['Heading1'],
            fontSize=24,
            textColor=colors.HexColor('#4263EB'),
            alignment=TA_CENTER,
            spaceAfter=6,
            fontName='DejaVuSans-Bold'
        )

        subtitle_style = ParagraphStyle(
            'ReportSubtitle',
            parent=styles['Normal'],
            fontSize=12,
            textColor=colors.HexColor('#64748b'),
            alignment=TA_CENTER,
            spaceAfter=4,
            fontName='DejaVuSans'
        )

        generated_style = ParagraphStyle(
            'GeneratedStyle',
            parent=styles['Normal'],
            fontSize=9,
            textColor=colors.HexColor('#94a3b8'),
            alignment=TA_CENTER,
            spaceAfter=16,
            fontName='DejaVuSans'
        )

        section_header_style = ParagraphStyle(
            'SectionHeader',
            parent=styles['Heading2'],
            fontSize=13,
            textColor=colors.HexColor('#1e293b'),
            spaceBefore=16,
            spaceAfter=10,
            fontName='DejaVuSans-Bold',
            borderWidth=0,
            borderPadding=0,
            borderColor=colors.HexColor('#e2e8f0')
        )

        normal_style = ParagraphStyle(
            'CustomNormal',
            parent=styles['Normal'],
            fontSize=9,
            textColor=colors.HexColor('#1e293b'),
            fontName='DejaVuSans'
        )

        # Format date range
        date_from = report_data.get('date_from', '')
        date_to = report_data.get('date_to', '')
        try:
            d_from = date.fromisoformat(date_from)
            d_to = date.fromisoformat(date_to)
            date_range_str = f"{d_from.strftime('%d %B %Y')} — {d_to.strftime('%d %B %Y')}"
        except:
            date_range_str = f"{date_from} — {date_to}"

        generated_str = f"Generated on {date.today().strftime('%d %B %Y')}"

        # ============ PAGE 1: Header, Summary & Category Breakdown ============

        # Header
        elements.append(Paragraph("FINANCE REPORT", title_style))
        elements.append(Paragraph(date_range_str, subtitle_style))
        elements.append(Paragraph(generated_str, generated_style))

        # Separator line
        separator_data = [[''] ]
        separator = Table(separator_data, colWidths=[page_width - 2 * margin])
        separator.setStyle(TableStyle([
            ('LINEBELOW', (0, 0), (-1, 0), 2, colors.HexColor('#4263EB')),
            ('TOPPADDING', (0, 0), (-1, 0), 0),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 8),
        ]))
        elements.append(separator)
        elements.append(Spacer(1, 8))

        # Summary Section
        elements.append(Paragraph("SUMMARY", section_header_style))

        total_income = report_data.get('total_income', 0)
        total_expenses = report_data.get('total_expenses', 0)
        net_savings = report_data.get('net_savings', 0)
        savings_rate = report_data.get('savings_rate', 0)
        income_count = report_data.get('income_count', 0)
        expense_count = report_data.get('expense_count', 0)

        def fmt_currency(val):
            return f"₹{val:,.0f}"

        summary_data = [
            ['Total Income', 'Total Expenses', 'Net Savings', 'Savings Rate'],
            [fmt_currency(total_income), fmt_currency(total_expenses), fmt_currency(net_savings), f"{savings_rate}%"]
        ]

        col_w = (page_width - 2 * margin) / 4
        summary_table = Table(summary_data, colWidths=[col_w] * 4)
        summary_table.setStyle(TableStyle([
            # Header row
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#f8fafc')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.HexColor('#64748b')),
            ('FONTNAME', (0, 0), (-1, 0), 'DejaVuSans'),
            ('FONTSIZE', (0, 0), (-1, 0), 8),
            ('ALIGN', (0, 0), (-1, 0), 'CENTER'),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 4),
            # Value row
            ('BACKGROUND', (0, 1), (-1, 1), colors.HexColor('#f8fafc')),
            ('FONTNAME', (0, 1), (-1, 1), 'DejaVuSans-Bold'),
            ('FONTSIZE', (0, 1), (-1, 1), 14),
            ('ALIGN', (0, 1), (-1, 1), 'CENTER'),
            ('TOPPADDING', (0, 1), (-1, 1), 4),
            ('BOTTOMPADDING', (0, 1), (-1, 1), 8),
            # Colors
            ('TEXTCOLOR', (0, 1), (0, 1), colors.HexColor('#16a34a')),  # Income green
            ('TEXTCOLOR', (1, 1), (1, 1), colors.HexColor('#dc2626')),  # Expense red
            ('TEXTCOLOR', (2, 1), (2, 1), colors.HexColor('#4263EB')),  # Savings blue
            ('TEXTCOLOR', (3, 1), (3, 1), colors.HexColor('#7c3aed')),  # Rate purple
            # Border
            ('BOX', (0, 0), (-1, -1), 1, colors.HexColor('#e2e8f0')),
            ('LINEBELOW', (0, 0), (-1, 0), 1, colors.HexColor('#e2e8f0')),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ]))
        elements.append(summary_table)
        elements.append(Spacer(1, 6))

        # Transaction count
        count_data = [[f"Transactions: {income_count + expense_count} total ({income_count} income, {expense_count} expense)"]]
        count_table = Table(count_data, colWidths=[page_width - 2 * margin])
        count_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#f8fafc')),
            ('BOX', (0, 0), (-1, -1), 1, colors.HexColor('#e2e8f0')),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, -1), 'DejaVuSans'),
            ('FONTSIZE', (0, 0), (-1, -1), 9),
            ('TEXTCOLOR', (0, 0), (-1, -1), colors.HexColor('#64748b')),
            ('TOPPADDING', (0, 0), (-1, -1), 6),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
        ]))
        elements.append(count_table)
        elements.append(Spacer(1, 16))

        # Spending by Category
        elements.append(Paragraph("SPENDING BY CATEGORY", section_header_style))

        category_breakdown = report_data.get('category_breakdown', [])
        if category_breakdown:
            cat_headers = ['Category', 'Transactions', '%', 'Amount']
            cat_data = [cat_headers]
            for cat in category_breakdown:
                cat_data.append([
                    cat.get('category_name', 'Uncategorized'),
                    str(cat.get('count', 0)),
                    f"{cat.get('percentage', 0)}%",
                    fmt_currency(cat.get('total', 0))
                ])
            cat_data.append(['TOTAL', str(sum(c.get('count', 0) for c in category_breakdown)),
                           '100%', fmt_currency(total_expenses)])

            cat_col_widths = [(page_width - 2 * margin) * 0.4,
                            (page_width - 2 * margin) * 0.2,
                            (page_width - 2 * margin) * 0.15,
                            (page_width - 2 * margin) * 0.25]

            cat_table = Table(cat_data, colWidths=cat_col_widths)
            cat_style = [
                # Header
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#4263EB')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
                ('FONTNAME', (0, 0), (-1, 0), 'DejaVuSans-Bold'),
                ('FONTSIZE', (0, 0), (-1, 0), 8),
                ('ALIGN', (0, 0), (0, -1), 'LEFT'),
                ('ALIGN', (1, 0), (2, -1), 'CENTER'),
                ('ALIGN', (3, 0), (3, -1), 'RIGHT'),
                ('BOTTOMPADDING', (0, 0), (-1, 0), 8),
                ('TOPPADDING', (0, 0), (-1, 0), 8),
                # Body
                ('FONTNAME', (0, 1), (-1, -2), 'DejaVuSans'),
                ('FONTSIZE', (0, 1), (-1, -1), 9),
                ('TOPPADDING', (0, 1), (-1, -1), 6),
                ('BOTTOMPADDING', (0, 1), (-1, -1), 6),
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                # Footer
                ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('#f1f5f9')),
                ('FONTNAME', (0, -1), (-1, -1), 'DejaVuSans-Bold'),
                ('LINEABOVE', (0, -1), (-1, -1), 2, colors.HexColor('#4263EB')),
                # Grid
                ('BOX', (0, 0), (-1, -1), 1, colors.HexColor('#e2e8f0')),
                ('LINEBELOW', (0, 0), (-1, -2), 1, colors.HexColor('#e2e8f0')),
            ]
            # Alternating row colors
            for i in range(1, len(cat_data) - 1):
                if i % 2 == 0:
                    cat_style.append(('BACKGROUND', (0, i), (-1, i), colors.HexColor('#f8fafc')))

            cat_table.setStyle(TableStyle(cat_style))
            elements.append(cat_table)

        # Footer for page 1
        elements.append(Spacer(1, 20))
        footer_style = ParagraphStyle(
            'FooterStyle',
            parent=styles['Normal'],
            fontSize=8,
            textColor=colors.HexColor('#94a3b8'),
            alignment=TA_CENTER,
            borderWidth=1,
            borderColor=colors.HexColor('#e2e8f0'),
            borderPadding=8,
            fontName='DejaVuSans'
        )
        elements.append(Paragraph("FinanceTrack — Personal Finance Manager", footer_style))

        # ============ PAGE 2+: All Transactions ============
        elements.append(PageBreak())

        elements.append(Paragraph("ALL TRANSACTIONS", section_header_style))

        transactions = report_data.get('transactions', [])
        if transactions:
            tx_headers = ['Date', 'Type', 'Category', 'Notes', 'Amount']
            tx_data = [tx_headers]
            for tx in transactions:
                tx_type = 'IN' if tx.get('type') == 'income' else 'EXP'
                tx_data.append([
                    tx.get('date', ''),
                    tx_type,
                    tx.get('category', ''),
                    tx.get('notes', '') or '-',
                    fmt_currency(tx.get('amount', 0))
                ])

            # Summary rows
            tx_data.append([f'Total Income ({income_count} transactions)', '', '', '', fmt_currency(total_income)])
            tx_data.append([f'Total Expenses ({expense_count} transactions)', '', '', '', fmt_currency(total_expenses)])
            tx_data.append(['Net Savings', '', '', '', fmt_currency(net_savings)])

            tx_col_widths = [(page_width - 2 * margin) * 0.14,
                           (page_width - 2 * margin) * 0.10,
                           (page_width - 2 * margin) * 0.18,
                           (page_width - 2 * margin) * 0.38,
                           (page_width - 2 * margin) * 0.20]

            tx_table = Table(tx_data, colWidths=tx_col_widths, repeatRows=1)
            tx_style = [
                # Header
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#4263EB')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
                ('FONTNAME', (0, 0), (-1, 0), 'DejaVuSans-Bold'),
                ('FONTSIZE', (0, 0), (-1, 0), 8),
                ('ALIGN', (0, 0), (0, -1), 'LEFT'),
                ('ALIGN', (1, 0), (1, -1), 'CENTER'),
                ('ALIGN', (2, 0), (2, -1), 'LEFT'),
                ('ALIGN', (3, 0), (3, -1), 'LEFT'),
                ('ALIGN', (4, 0), (4, -1), 'RIGHT'),
                ('BOTTOMPADDING', (0, 0), (-1, 0), 8),
                ('TOPPADDING', (0, 0), (-1, 0), 8),
                # Body
                ('FONTNAME', (0, 1), (-1, -4), 'DejaVuSans'),
                ('FONTSIZE', (0, 1), (-1, -1), 8),
                ('TOPPADDING', (0, 1), (-1, -1), 5),
                ('BOTTOMPADDING', (0, 1), (-1, -1), 5),
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                # Footer summary rows
                ('BACKGROUND', (0, -3), (-1, -3), colors.HexColor('#f1f5f9')),
                ('BACKGROUND', (0, -2), (-1, -2), colors.HexColor('#f1f5f9')),
                ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('#f1f5f9')),
                ('FONTNAME', (0, -3), (-1, -1), 'DejaVuSans-Bold'),
                ('LINEABOVE', (0, -3), (-1, -3), 2, colors.HexColor('#4263EB')),
                # Grid
                ('BOX', (0, 0), (-1, -1), 1, colors.HexColor('#e2e8f0')),
                ('LINEBELOW', (0, 0), (-1, -4), 1, colors.HexColor('#e2e8f0')),
            ]

            # Type badge colors & alternating rows
            for i in range(1, len(tx_data) - 3):
                if i % 2 == 0:
                    tx_style.append(('BACKGROUND', (0, i), (-1, i), colors.HexColor('#f8fafc')))
                # Color the type column
                tx_type_val = tx_data[i][1]
                if tx_type_val == 'IN':
                    tx_style.append(('TEXTCOLOR', (1, i), (1, i), colors.HexColor('#16a34a')))
                else:
                    tx_style.append(('TEXTCOLOR', (1, i), (1, i), colors.HexColor('#dc2626')))

            tx_table.setStyle(TableStyle(tx_style))
            elements.append(tx_table)

        # Footer for transactions page
        elements.append(Spacer(1, 20))
        elements.append(Paragraph("FinanceTrack — Personal Finance Manager", footer_style))

        # ============ LAST PAGE: Payment Method Breakdown ============
        elements.append(PageBreak())

        elements.append(Paragraph("PAYMENT METHOD BREAKDOWN", section_header_style))

        payment_methods = report_data.get('payment_methods', [])
        if payment_methods:
            pm_headers = ['Payment Method', 'Transactions', '%', 'Amount']
            pm_data = [pm_headers]
            for pm in payment_methods:
                pm_data.append([
                    pm.get('method', ''),
                    str(pm.get('count', 0)),
                    f"{pm.get('percentage', 0)}%",
                    fmt_currency(pm.get('total', 0))
                ])
            pm_data.append(['TOTAL', str(sum(p.get('count', 0) for p in payment_methods)),
                          '100%', fmt_currency(total_expenses)])

            pm_col_widths = [(page_width - 2 * margin) * 0.4,
                           (page_width - 2 * margin) * 0.2,
                           (page_width - 2 * margin) * 0.15,
                           (page_width - 2 * margin) * 0.25]

            pm_table = Table(pm_data, colWidths=pm_col_widths)
            pm_style = [
                # Header
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#4263EB')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
                ('FONTNAME', (0, 0), (-1, 0), 'DejaVuSans-Bold'),
                ('FONTSIZE', (0, 0), (-1, 0), 8),
                ('ALIGN', (0, 0), (0, -1), 'LEFT'),
                ('ALIGN', (1, 0), (2, -1), 'CENTER'),
                ('ALIGN', (3, 0), (3, -1), 'RIGHT'),
                ('BOTTOMPADDING', (0, 0), (-1, 0), 8),
                ('TOPPADDING', (0, 0), (-1, 0), 8),
                # Body
                ('FONTNAME', (0, 1), (-1, -2), 'DejaVuSans'),
                ('FONTSIZE', (0, 1), (-1, -1), 9),
                ('TOPPADDING', (0, 1), (-1, -1), 6),
                ('BOTTOMPADDING', (0, 1), (-1, -1), 6),
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                # Footer
                ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('#f1f5f9')),
                ('FONTNAME', (0, -1), (-1, -1), 'DejaVuSans-Bold'),
                ('LINEABOVE', (0, -1), (-1, -1), 2, colors.HexColor('#4263EB')),
                # Grid
                ('BOX', (0, 0), (-1, -1), 1, colors.HexColor('#e2e8f0')),
                ('LINEBELOW', (0, 0), (-1, -2), 1, colors.HexColor('#e2e8f0')),
            ]
            for i in range(1, len(pm_data) - 1):
                if i % 2 == 0:
                    pm_style.append(('BACKGROUND', (0, i), (-1, i), colors.HexColor('#f8fafc')))

            pm_table.setStyle(TableStyle(pm_style))
            elements.append(pm_table)

        # Disclaimer
        elements.append(Spacer(1, 30))
        disclaimer_style = ParagraphStyle(
            'DisclaimerStyle',
            parent=styles['Normal'],
            fontSize=8,
            textColor=colors.HexColor('#94a3b8'),
            alignment=TA_CENTER,
            fontName='DejaVuSans'
        )
        elements.append(Paragraph("This report was generated automatically. Please verify all transactions.", disclaimer_style))

        # Build PDF
        doc.build(elements)
        output.seek(0)
        return output.getvalue()
    except ImportError:
        return None
